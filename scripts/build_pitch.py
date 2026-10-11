import shutil
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.opc.packuri import PackURI

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "syntax-template.pptx"
OUTPUT = ROOT / "syntax-презентация.pptx"
SHOTS = ROOT / "design" / "app-screenshots"


def set_text(slide, idx, text):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx and ph.has_text_frame:
            ph.text_frame.text = text
            return ph
    return None


def insert_picture(slide, idx, path, crop_focus_x=0.5):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx and ph.placeholder_format.type == 18 and hasattr(ph, "insert_picture"):
            picture = ph.insert_picture(str(path))
            return picture
    return None


def build():
    source = TEMPLATE if TEMPLATE.exists() else ROOT / "syntax-template-base.pptx"
    if not source.exists():
        source = OUTPUT
    if source != OUTPUT:
        shutil.copyfile(source, OUTPUT)
    prs = Presentation(str(OUTPUT))
    slides = list(prs.slides)
    if len(slides) == 11:
        slides = [None, *slides, None]
    if len(slides) > 13:
        keep = [1, 7, 11, 24, 15, 16, 21, 17, 26, 29, 25, 13, 9]
        entries = list(prs.slides._sldIdLst)
        for entry in entries:
            prs.slides._sldIdLst.remove(entry)
        for number in keep:
            prs.slides._sldIdLst.append(entries[number - 1])
        slides = list(prs.slides)

    title = slides[1]
    set_text(title, 0, "syntax")
    set_text(title, 12, "Цифровая платформа-агрегатор ИТ-вакансий с верифицированным профилем достижений ПВН\nПлатформа обратного найма · спецтрек ПВН 2026")

    short = slides[2]
    set_text(short, 0, "КОРОТКО О РЕШЕНИИ")
    set_text(short, 49, "Как работает подбор")
    set_text(short, 38, "Работодатель ищет не по резюме, а по верифицированным категориям: специализация × грейд, подтверждённые адаптивным тестом.\n\nАнонимные карточки кандидатов с объяснимым скорингом: тесты 40% · достижения ПВН 30% · регулярные задания 20% · полнота профиля 10%.\n\nКонтакты кандидата раскрываются только после принятия приглашения — обе стороны видят зарплатную вилку до первого контакта.")
    set_text(short, 51, "Для кого платформа")
    set_text(short, 42, "Обратный найм: инициатива у работодателя, кандидат не откликается вслепую.\n\nКонкурентное преимущество для участников ПВН: спортивные достижения конвертируются в карьерный бонус и поднимают в ранжировании.\n\nЦелевые сегменты: ИТ-компании, ищущие точный подбор, и разработчики, уставшие от откликов без ответа.")

    problem = slides[3]
    set_text(problem, 0, "Задача подбора")
    set_text(problem, 26, "Как сегодня\n\nРаботодатель вручную фильтрует сотни нерелевантных откликов силами дорогих технических специалистов")
    set_text(problem, 31, "Почему так\n\nРезюме фиксирует стек, но не подтверждает компетентность: отбор сводится к совпадению ключевых слов")
    set_text(problem, 32, "Insight: syntax\n\nРаботодатель выбирает из верифицированных категорий, зарплата видна до контакта, контакты — после взаимного согласия")

    matching = slides[4]
    set_text(matching, 0, "Механика подбора")
    set_text(matching, 15, "Категория = специализация × грейд — по итогам опроса и тестирования, а не самоописанного резюме")
    set_text(matching, 16, "Анонимные карточки: кандидат #A17F без имени и контактов — только грейд, стек, скоринг и достижения")
    set_text(matching, 17, "Объяснимый скоринг: score = Σ(компонент × вес), при отсутствии истории ПВН веса перенормируются")
    set_text(matching, 18, "Ранжирование Score × MatchRelevance: semantic 70% + stack 30% при актуальном embedding, иначе lexical fallback")
    set_text(matching, 19, "Приглашение с обязательной вилкой «от–до» в рублях: 422 при инверсии, контакты скрыты до статуса accepted")

    testing = slides[5]
    set_text(testing, 0, "Тестирование")
    set_text(testing, 49, "1")
    set_text(testing, 50, "2")
    set_text(testing, 51, "3")
    set_text(testing, 52, "4")
    set_text(testing, 37, "Банк концептов")
    set_text(testing, 38, "40 концептов, 5 доменов × 3 грейда, якоря сложности: сложность заякорена параметрами, а не текстом задачи")
    set_text(testing, 39, "Каскад генерации")
    set_text(testing, 40, "LLM → fallback-пул из 240 задач → детерминированная генерация по seed «кандидат:сессия»: формулировка уникальна каждому, опции перемешаны")
    set_text(testing, 41, "Вердикт")
    set_text(testing, 42, "5 задач на грейд: 4–5 верных → предложение выше, 3 → подтверждён, 0–2 → рекомендация ниже; грейд не понижается принудительно")
    set_text(testing, 43, "Кулдаун грейда")
    set_text(testing, 44, "Смена грейда — решение кандидата, не чаще раза в 3 месяца (конфигурируется); регулярные задания поддерживают профиль в актуальном состоянии")

    validation = slides[6]
    set_text(validation, 0, "Валидация")
    set_text(validation, 21, "236 backend-тестов + 10 pgvector")
    set_text(validation, 18, "Последний прогон: Docker PostgreSQL/pgvector с rollback временных записей")
    set_text(validation, 22, "Self-reported метрики")
    set_text(validation, 23, "Прежние AUC, test-retest и Precision@5 не прошли независимую проверку")
    set_text(validation, 24, "UX-аудит: текстовая модель")
    set_text(validation, 25, "Сценарий и изображение принимаются, но текущая модель не выполняет визуальный анализ; реальный PVN/ATS не настроен")

    pvn = slides[7]
    set_text(pvn, 0, "Интеграция ПВН")
    set_text(pvn, 49, "1")
    set_text(pvn, 50, "2")
    set_text(pvn, 51, "3")
    set_text(pvn, 37, "Достижения → скоринг")
    set_text(pvn, 38, "Олимпиады, хакатоны, проекты ПВН дают компонент 30% в скоринге и поднимают кандидата в ранжировании категории")
    set_text(pvn, 39, "Нет истории ПВН")
    set_text(pvn, 40, "Обязательный кейс из ТЗ: компонент исключается, веса перенормируются, карточка получает явную пометку — достижения бонус, а не фильтр")
    set_text(pvn, 41, "Проверка OIDC-токена")
    set_text(pvn, 42, "OIDC/JWKS exchange path реализован с проверкой issuer + subject; реальный провайдер, реестр и SSO в окружении не настроены")

    demo_match = slides[8]
    set_text(demo_match, 0, "Демо: работодатель находит кандидата")
    insert_picture(demo_match, 14, SHOTS / "02-employer-match.png")
    set_text(demo_match, 15, "Категории и анонимные карточки: Backend · Middle · Python — 47 кандидатов")
    set_text(demo_match, 16, "Объяснимый скоринг: разбор по компонентам с весами в каждой карточке")
    set_text(demo_match, 17, "Приглашение с зарплатной вилкой до раскрытия контактов")

    demo_flow = slides[9]
    set_text(demo_flow, 0, "Путь до контакта")
    insert_picture(demo_flow, 14, SHOTS / "03-invitation-modal.png")
    insert_picture(demo_flow, 15, SHOTS / "04-candidate-invitations.png")
    insert_picture(demo_flow, 16, SHOTS / "05-testing-session.png")
    insert_picture(demo_flow, 17, SHOTS / "06-candidate-profile.png")
    set_text(demo_flow, 18, "Работодатель отправляет приглашение с вилкой 200 000 – 260 000 ₽")
    set_text(demo_flow, 19, "Кандидат видит условия заранее и принимает приглашение")
    set_text(demo_flow, 20, "Адаптивный тест: 5 задач, вердикт по грейду")
    set_text(demo_flow, 21, "Анонимная карточка и полные данные в одном профиле")

    journey = slides[10]
    set_text(journey, 0, "Путь кандидата")
    set_text(journey, 26, "Регистрация")
    set_text(journey, 27, "Email с подтверждением и согласие 152-ФЗ; PVN ID — необязательное локальное поле")
    set_text(journey, 28, "Опрос и тест")
    set_text(journey, 29, "Специализация, заявленный грейд, 5 адаптивных задач")
    set_text(journey, 30, "Категория")
    set_text(journey, 31, "Автоматическое попадание в категорию работодателя")
    set_text(journey, 32, "Приглашение")
    set_text(journey, 33, "Входящие офферы с зарплатной вилкой и описанием")
    set_text(journey, 34, "Контакт")
    set_text(journey, 35, "Принятие → контакты раскрыты обеим сторонам")

    architecture = slides[11]
    set_text(architecture, 0, "Архитектура")
    set_text(architecture, 1, "Монолит с явным разделением слоёв: api / core / services / models / auth.\n\nFrontend: React SPA (Vite + TypeScript + Tailwind), nginx в docker.\nХранение: PostgreSQL 16, миграции Alembic.\n\nLLM-провайдер опционален: при отсутствии работает fallback-пул задач и детерминированная генерация. Embedding-подбор деградирует до lexical fallback.\n\nКонтейнеризация: docker-compose (api :8000, db, web :3000). Последний прогон: 236 backend-тестов и 10 PostgreSQL/pgvector тестов с rollback. Внешние PVN/ATS не подтверждены.")

    team = slides[12]
    if team is not None:
        set_text(team, 0, "Команда")
        team_roles = ["Продукт и сценарий", "Backend и интеграции", "Frontend и UX", "Тестирование и данные", "Демо и документация"]
        role_index = 0
        for shape in team.shapes:
            if not shape.has_text_frame or shape.text.strip() != "Имя Фамилия":
                continue
            shape.text_frame.text = team_roles[role_index]
            role_index += 1

    set_text(title, 12, "Платформа обратного найма ИТ-специалистов\nСпецтрек ПВН 2026")
    set_text(short, 38, "Категории по специализации и грейду; тестирование даёт сигнал о навыках.\n\nАнонимные карточки и объяснимый скоринг. Отсутствующие ПВН-достижения не получают нулевой вклад.\n\nЗарплата видна заранее. Контакты раскрываются после принятия приглашения.")
    set_text(short, 42, "Работодатель приглашает кандидата по понятным сигналам профиля.\n\nДостижения ПВН поддерживаются как локальные записи; реестр пока не подключён.\n\nЦелевой пользователь: ИТ-компания и специалист, которому важны условия до первого контакта.")
    set_text(pvn, 38, "Локальные записи достижений дают компонент 30% в скоринге. Это не подтверждённый обмен с реестром ПВН.")
    set_text(testing, 40, "LLM → пул из 240 задач → seed-генерация. Формулировки и порядок ответов варьируются; устойчивость к утечке требует проверки.")
    for slide_number, slide in enumerate(slides, start=1):
        if slide is None:
            continue
        for shape in slide.shapes:
            if getattr(shape, "has_chart", False):
                shape._element.getparent().remove(shape._element)
                continue
            if not shape.has_text_frame:
                continue
            if shape.is_placeholder and shape.placeholder_format.type == 13:
                shape.text_frame.text = str(slide_number - 1)
            if shape.is_placeholder and slide_number == 6 and shape.placeholder_format.idx in (49, 50, 51, 52):
                shape.height = max(shape.height, 450000)
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    dark_background = slide_number in (1, 2) or (slide_number in (3, 4, 6, 7, 9, 10, 12) and shape.is_placeholder and shape.placeholder_format.idx in (0, 4))
                    run.font.color.rgb = RGBColor(255, 255, 255) if dark_background else RGBColor(40, 24, 55)
    slide_parts = [part for part in prs.part.package.iter_parts() if part.partname.startswith("/ppt/slides/") and part.partname.endswith(".xml")]
    for number, part in enumerate(slide_parts, start=1):
        part._partname = PackURI(f"/ppt/slides/rebuilt{number}.xml")
    prs.save(str(OUTPUT))
    print(f"built {OUTPUT} with {len(prs.slides)} slides")


if __name__ == "__main__":
    build()
